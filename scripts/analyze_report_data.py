"""Reproduce descriptive pilot results from the retained questionnaire workbook.

This path uses recorded completion, not unavailable physical trial telemetry.
The stricter reachy_lab2.analysis path remains available for telemetry-linked data.
"""
import argparse
import csv
import hashlib
import json
import statistics as stats
from pathlib import Path

DIMENSIONS = {"sociability": "soc", "animacy": "ani", "agency": "age", "disturbance": "dis"}
TERMS = ["Warm", "Likeable", "Trustworthy", "Friendly", "Alive", "Natural", "Real", "Human-like",
         "Self-reliant", "Rational", "Intentional", "Intelligent", "Creepy", "Scary", "Uncanny", "Weird"]
CODES = {("P01", "A"): "Robotic or mechanical impression", ("P01", "B"): "Increased liveliness or interactivity",
         ("P02", "A"): "Little or no perceived difference", ("P02", "B"): "Robotic or mechanical impression",
         ("P03", "A"): "Robotic or mechanical impression", ("P03", "B"): "Increased liveliness or interactivity",
         ("P04", "A"): "Little or no perceived difference", ("P04", "B"): "Little or no perceived difference"}


def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path, records):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)


def rating(value):
    if value in (None, "", "NA"):
        return None
    number = int(value)
    if number != float(value) or not 1 <= number <= 7:
        raise ValueError(f"Invalid rating: {value!r}")
    return number


def normalized_text(value):
    return str(value).replace("’", "'").replace("‘", "'")


def analyze(root, out):
    from openpyxl import load_workbook
    root, out = Path(root).resolve(), Path(out).resolve()
    raw = root / "study/data/raw"
    workbook = raw / "De-identified raw data.xlsx"
    sheet = load_workbook(workbook, data_only=True)["Trial Data"]
    values = list(sheet.iter_rows(values_only=True))
    source = [dict(zip(values[0], row)) for row in values[1:] if row[0] is not None]
    participants = {r["participant_id"]: r for r in read_csv(raw / "participants.csv")}
    csv_rows = {(r["participant_id"], r["condition"]): r for r in read_csv(raw / "responses.csv")}
    deviations = {(r["participant_id"], r["condition"]): r["deviation"] for r in read_csv(raw / "deviations.csv")}
    conditions = json.loads((root / "config/conditions.json").read_text())
    protocol = json.loads((root / "study/protocol.json").read_text())
    if conditions["conditions"] != {"A": 10.0, "B": 30.0}:
        raise ValueError("Report stimulus settings changed; revise report before regenerating")
    ready, seen = [], set()
    for r in source:
        pid, condition = r["Participant ID"], r["Condition"]
        key = pid, condition
        if key in seen or condition not in "AB":
            raise ValueError("Duplicate or invalid participant condition")
        seen.add(key)
        p, c = participants[pid], csv_rows[key]
        order = r["Order"]
        if not order == p["actual_order"] == p["assigned_order"] == protocol["assignments"][pid]:
            raise ValueError("Condition order mismatch")
        if r["Greeting #"] != order.index(condition) + 1:
            raise ValueError("Greeting position mismatch")
        row = dict(participant_id=pid, gender=r["Gender"], age_band=p["age_band"],
                   robot_experience=p["robot_experience"], reachy_experience=p["reachy_experience"],
                   agreement=p["agreement"], session_date=p["date"], operator=p["operator"],
                   condition=condition, antenna_amplitude_deg=conditions["conditions"][condition],
                   order=order, greeting_number=r["Greeting #"], study_version=protocol["study_version"],
                   survey_version=protocol["survey_version"])
        for term, name in zip(TERMS, [f"{prefix}_{i}" for prefix in DIMENSIONS.values() for i in range(1, 5)]):
            row[name] = rating(r[term])
            if row[name] != rating(c[name]):
                raise ValueError(f"Workbook/CSV mismatch: {key}/{name}")
        for dimension, prefix in DIMENSIONS.items():
            items = [row[f"{prefix}_{i}"] for i in range(1, 5)]
            row[dimension] = None if None in items else stats.mean(items)
        for measure, label in [("valence", "Valence"), ("arousal", "Arousal"), ("movement_size", "Movement size")]:
            row[measure] = rating(r[label])
            if row[measure] != rating(c[measure]):
                raise ValueError("Workbook/CSV rating mismatch")
        row["qualitative_response"] = r["Open-ended response"]
        if normalized_text(row["qualitative_response"]) != normalized_text(c["qualitative_response"]):
            raise ValueError("Workbook/CSV qualitative mismatch")
        row["completion"] = None if r["Completion"] in (None, "", "NA") else int(r["Completion"])
        if row["completion"] not in (0, 1, None) or str(row["completion"]) != c["completion"]:
            raise ValueError("Invalid or mismatched completion")
        row["completion_source"] = "recorded worksheet outcome; physical logs unavailable"
        row["failure_interruption_reason"] = r["Failure / interruption reason"] or ""
        row["protocol_deviation"] = deviations[key]
        ready.append(row)
    if seen != set(csv_rows) or len(participants) != 4 or len(ready) != 8:
        raise ValueError("Participant accounting changed; revise report")
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / "analysis_ready.csv", ready)
    summaries, pairs = [], []
    measures = [*DIMENSIONS, "valence", "arousal", "movement_size"]
    for measure in measures:
        paired = []
        for pid in sorted(participants):
            by = {r["condition"]: r for r in ready if r["participant_id"] == pid}
            a, b = by["A"][measure], by["B"][measure]
            diff = None if a is None or b is None else b - a
            paired.append(diff)
            pairs.append(dict(participant_id=pid, measure=measure, A=a, B=b, B_minus_A=diff))
        for condition in "AB":
            vals = [r[measure] for r in ready if r["condition"] == condition and r[measure] is not None]
            summaries.append(dict(measure=measure, condition=condition, n=len(vals), mean=stats.mean(vals),
                                  sample_sd=stats.stdev(vals), median=stats.median(vals)))
        diffs = [d for d in paired if d is not None]
        summaries.append(dict(measure=measure, condition="B_minus_A", n=len(diffs), mean=stats.mean(diffs),
                              sample_sd=stats.stdev(diffs), median=stats.median(diffs)))
    write_csv(out / "condition_summary.csv", summaries)
    write_csv(out / "paired_differences.csv", pairs)
    objective = []
    for condition in "AB":
        vals = [r["completion"] for r in ready if r["condition"] == condition and r["completion"] is not None]
        objective.append(dict(condition=condition, recorded_trials=len(vals), completed=sum(vals),
                              completion_proportion=stats.mean(vals), telemetry_verified="unavailable"))
    write_csv(out / "completion_summary.csv", objective)
    qualitative = [dict(participant_id=r["participant_id"], condition=r["condition"],
                        qualitative_response=r["qualitative_response"], category=CODES[(r["participant_id"], r["condition"])]
                        ) for r in ready]
    write_csv(out / "qualitative_coding.csv", qualitative)
    missing = {field: sum(r[field] is None for r in ready) for field in
               [f"{p}_{i}" for p in DIMENSIONS.values() for i in range(1, 5)] + measures + ["completion"]}
    audit = dict(participants=len(participants), observations=len(ready), missing=missing,
                 comparison="B minus A", primary_hries="sociability", physical_trial_logs="unavailable in submission",
                 workbook_csv_match=True, text_comparison="apostrophe typography normalized for comparison only",
                 date_source="participants.csv; workbook date cells blank",
                 sources={str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in
                          [workbook, raw / "responses.csv", raw / "participants.csv", raw / "deviations.csv",
                           root / "config/conditions.json", root / "study/protocol.json"]})
    (out / "audit.json").write_text(json.dumps(audit, indent=2) + "\n")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    colors = ["#0072B2", "#D55E00", "#009E73", "#CC79A7"]
    def paired_plot(names, labels, filename, title):
        fig, axes = plt.subplots(1, len(names), figsize=(7.1, 3.35), sharey=True, squeeze=False)
        for ax, measure, label in zip(axes.flat, names, labels):
            for pid, color in zip(sorted(participants), colors):
                pair = next(p for p in pairs if p["participant_id"] == pid and p["measure"] == measure)
                ax.plot([0, 1], [pair["A"], pair["B"]], "o-", color=color, label=pid, lw=1.7, markersize=5)
            ax.set(title=label, xticks=[0, 1], xticklabels=["A\n10°", "B\n30°"], ylim=(.7, 7.3), yticks=range(1, 8), xlim=(-.15, 1.15))
            ax.grid(axis="y", color=".9")
            ax.spines[["top", "right"]].set_visible(False)
        axes[0, 0].set_ylabel("Rating or dimension mean (1–7)")
        handles, labels_ = axes[0, 0].get_legend_handles_labels()
        fig.legend(handles, labels_, loc="lower center", ncol=4, frameon=False)
        fig.suptitle(title, fontsize=12)
        fig.tight_layout(rect=(0, .12, 1, .94))
        for suffix in ("png", "svg"):
            fig.savefig(out / f"{filename}.{suffix}", dpi=240)
        plt.close(fig)
    paired_plot(list(DIMENSIONS), ["Sociability", "Animacy", "Agency", "Disturbance"], "paired_hries", "HRIES dimensions for all four participants")
    paired_plot(["valence", "arousal", "movement_size"], ["Valence", "Arousal", "Movement size"], "paired_affect", "Perceived affect and movement size")
    return summaries


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--out", type=Path, default=Path("study/final_submission/analysis"))
    args = parser.parse_args()
    for row in analyze(args.root, args.out):
        print(f"{row['measure']:14s} {row['condition']:10s} n={row['n']} mean={row['mean']:.4f} SD={row['sample_sd']:.4f}")
